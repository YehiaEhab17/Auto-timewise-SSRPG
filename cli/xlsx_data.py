from pathlib import Path

from core.classes import LOCATION_NAMES
from core.location_data import load_from_json, save_json

# display name -> internal name (reverse of LOCATION_NAMES)
LOCATIONS = {display: internal for internal, display in LOCATION_NAMES.items()}


def load_from_xlsx(sheet_path):
    import openpyxl

    LOCATION_VALUES: dict[str, float] = {}

    ROWS = [55, 64, 73, 82, 91, 100, 109, 118]

    wb = openpyxl.load_workbook(sheet_path, data_only=True)
    ws = wb["Chest Rates"]

    for row in ROWS:
        name = str(ws.cell(row, 2).value)
        values = [float(ws.cell(row + 7, col).value) for col in range(4, 20)]

        loc_id = LOCATIONS[name]
        star = 5

        for v in values:
            LOCATION_VALUES[loc_id + str(star)] = float(v)
            star += 1

    return LOCATION_VALUES


def get_location_values():
    script_location = Path(__file__).parent.parent
    sheet_path = script_location / "chest_rates.xlsx"
    json_path = script_location / "chest_rates.json"

    if sheet_path.exists():
        try:
            data = load_from_xlsx(sheet_path)
            save_json(json_path, data)
            return data
        except ImportError:
            print("openpyxl not found, try installing")

    else:
        print("chest_rates.xlsx sheet not found")

    print("could not use xlsx sheet, using cache")
    return load_from_json(json_path)
